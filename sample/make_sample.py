"""
Generates a synthetic business-planning workbook for local smoke testing.
All values and metadata are fictional and contain no customer records.

Run: python sample/make_sample.py
"""
import openpyxl

wb = openpyxl.Workbook()
wb.properties.creator = "Business Planning Example"
wb.properties.lastModifiedBy = "Business Planning Example"
wb.properties.title = "Synthetic Business Planning Workbook"
wb.properties.description = "Fictional sample data for workbook business-logic reverse engineering."

inputs = wb.active
inputs.title = "Inputs"
inputs["A1"] = "Period"
inputs["B1"] = "Units"
inputs["C1"] = "Unit Price"
inputs["D1"] = "Unit Cost"
data = [
    ("P1", 1000, 100, 40),
    ("P2", 1100, 100, 42),
    ("P3", 1050, 102, 41),
]
for i, (period, units, unit_price, unit_cost) in enumerate(data, start=2):
    inputs[f"A{i}"] = period
    inputs[f"B{i}"] = units
    inputs[f"C{i}"] = unit_price
    inputs[f"D{i}"] = unit_cost

plan = wb.create_sheet("Plan")
plan["A1"] = "Period"
plan["B1"] = "Planned Value"
plan["C1"] = "Planned Cost"
plan["D1"] = "Contribution"
plan["E1"] = "Contribution Rate"
for i in range(2, 5):
    plan[f"A{i}"] = f"=Inputs!A{i}"
    plan[f"B{i}"] = f"=Inputs!B{i}*Inputs!C{i}"
    plan[f"C{i}"] = f"=Inputs!B{i}*Inputs!D{i}"
    plan[f"D{i}"] = f"=B{i}-C{i}"
    plan[f"E{i}"] = f"=D{i}/B{i}"

summary = wb.create_sheet("Summary")
summary["A1"] = "Total Contribution"
summary["B1"] = "=SUM(Plan!D2:D4)"
summary["A2"] = "Avg Contribution Rate"
summary["B2"] = "=AVERAGE(Plan!E2:E4)"
summary["B2"].comment = openpyxl.comments.Comment(
    "The planning team reviews this synthetic KPI each period.", "SME"
)

wb.save("sample/business_planning_sample.xlsx")
print("Wrote sample/business_planning_sample.xlsx")

"""
excel_parser.py
----------------
Ingests an Excel workbook and reverse-engineers its structure:
  - worksheets, cells, formulas, static inputs
  - a dependency graph (which cells feed which cells)
  - per-sheet input/output classification (heuristic)

This is the raw "extraction" layer that Foundry-IQ-equivalent knowledge base
and the LLM agent consume downstream.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Collection, Dict, List, Optional

import networkx as nx
import openpyxl
from formulas import Parser
from formulas.tokens.operand import Range
from openpyxl.utils import coordinate_to_tuple, get_column_letter

MAX_EXPANDED_RANGE_CELLS = 10_000


@dataclass
class CellInfo:
    sheet: str
    coord: str
    value: Optional[object] = None
    formula: Optional[str] = None
    is_input: bool = False
    is_output: bool = False
    comment: Optional[str] = None

    @property
    def node_id(self) -> str:
        return f"{self.sheet}!{self.coord}"


@dataclass
class WorkbookModel:
    path: str
    sheets: List[str] = field(default_factory=list)
    cells: Dict[str, CellInfo] = field(default_factory=dict)
    graph: nx.DiGraph = field(default_factory=nx.DiGraph)
    authors: List[str] = field(default_factory=list)


def _extract_refs(
    formula: str,
    current_sheet: str,
    known_cells: Optional[Collection[str]] = None,
) -> List[str]:
    """Extract and expand cell references from a parsed Excel formula AST."""
    tokens, _ = Parser().ast(formula)
    refs: List[str] = []
    seen = set()

    for token in tokens:
        if not isinstance(token, Range):
            continue

        attrs = token.attr
        sheet = attrs.get("sheet") or current_sheet
        row_start = int(attrs["r1"])
        row_end = int(attrs["r2"])
        col_start = int(attrs["n1"])
        col_end = int(attrs["n2"])
        range_size = (row_end - row_start + 1) * (col_end - col_start + 1)

        if range_size > MAX_EXPANDED_RANGE_CELLS:
            candidates = []
            prefix = f"{sheet}!"
            for node_id in known_cells or ():
                if not node_id.startswith(prefix):
                    continue
                row, column = coordinate_to_tuple(node_id[len(prefix):])
                if row_start <= row <= row_end and col_start <= column <= col_end:
                    candidates.append(node_id)
        else:
            candidates = (
                f"{sheet}!{get_column_letter(column)}{row}"
                for row in range(row_start, row_end + 1)
                for column in range(col_start, col_end + 1)
            )

        for node_id in candidates:
            if node_id not in seen:
                seen.add(node_id)
                refs.append(node_id)

    return refs


def parse_workbook(path: str) -> WorkbookModel:
    wb = openpyxl.load_workbook(path, data_only=False)
    model = WorkbookModel(path=path, sheets=wb.sheetnames)

    # Work-IQ-equivalent: pull lightweight "who touched this" metadata.
    props = wb.properties
    for author in filter(None, [props.creator, props.lastModifiedBy]):
        if author not in model.authors:
            model.authors.append(author)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                coord = f"{get_column_letter(cell.column)}{cell.row}"
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                info = CellInfo(
                    sheet=sheet_name,
                    coord=coord,
                    value=None if is_formula else cell.value,
                    formula=cell.value if is_formula else None,
                    comment=cell.comment.text if cell.comment else None,
                )
                model.cells[info.node_id] = info
                model.graph.add_node(info.node_id)

    # Second pass: wire up dependency edges now that every node exists.
    for node_id, info in model.cells.items():
        if not info.formula:
            continue
        for ref in _extract_refs(info.formula, info.sheet, model.cells):
            model.graph.add_edge(ref, node_id)  # ref -> node_id (feeds into)

    # Heuristic input/output classification from graph topology.
    for node_id in model.graph.nodes:
        in_deg = model.graph.in_degree(node_id)
        out_deg = model.graph.out_degree(node_id)
        info = model.cells.get(node_id)
        if info is None:
            continue
        info.is_input = in_deg == 0 and info.formula is None
        info.is_output = out_deg == 0 and info.formula is not None

    return model


def summarize(model: WorkbookModel) -> str:
    lines = [f"Workbook: {model.path}", f"Sheets: {', '.join(model.sheets)}"]
    if model.authors:
        lines.append(f"Known authors (Work-IQ context): {', '.join(model.authors)}")
    n_formulas = sum(1 for c in model.cells.values() if c.formula)
    n_inputs = sum(1 for c in model.cells.values() if c.is_input)
    n_outputs = sum(1 for c in model.cells.values() if c.is_output)
    lines.append(f"Cells: {len(model.cells)} | Formulas: {n_formulas} | Inputs: {n_inputs} | Outputs: {n_outputs}")
    return "\n".join(lines)
